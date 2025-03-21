# import pandas as pd
# import openpyxl

# def read_excel_with_formulas_all_sheets(file_path):
#     """
#     Reads all sheets of an Excel file, including formulas, into a dictionary of DataFrames.

#     Args:
#         file_path (str): Path to the Excel file.

#     Returns:
#         dict: A dictionary where keys are sheet names and values are DataFrames containing data and formulas.
#     """
#     try:
#         wb = openpyxl.load_workbook(file_path, data_only=False) # Ensure data_only is False.
#         sheet_names = wb.sheetnames
#         all_dfs = {}

#         for sheet_name in sheet_names:
#             ws = wb[sheet_name]
#             df = pd.read_excel(file_path, sheet_name=sheet_name)
#             df['formula'] = None

#             for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column), start=1):
#                 for col_idx, cell in enumerate(row, start=1):
#                     cell.internal_value
#                     cell.col_idx
#                     print(f"Cell: {cell.coordinate}, Value: {cell.value}, Formula: {cell.formula if hasattr(cell, 'formula') else 'No Formula'}")
#                     if hasattr(cell, 'formula'): # Check if cell has a formula.
#                         df.loc[row_idx - 1, df.columns[col_idx - 1]] = f"{cell.value}" #replace value with openpyxl value in case of formula.
#                         df.loc[row_idx - 1, 'formula'] = cell.internal_value

#             all_dfs[sheet_name] = df

#         return all_dfs

#     except FileNotFoundError:
#         print(f"Error: File not found at {file_path}")
#         return None
#     except Exception as e:
#         print(f"An error occurred: {e}")
#         return None

# # Example usage:
# file_path = r"C:\Users\gaurav.j.choubey\Desktop\project\gtic-7-2025\excelsheet_poc\excel_genai_poc\examplemacro1.xlsm"  # Replace with your file path

# result_dfs = read_excel_with_formulas_all_sheets(file_path)

# if result_dfs:
#     for sheet_name, df in result_dfs.items():
#         print(f"Sheet: {sheet_name}")
#         print(df)
#         print("-" * 40)


import pandas as pd
import openpyxl
from openpyxl.utils import range_boundaries

def read_excel_with_formulas_all_sheets(file_path):
    """
    Reads all sheets of an Excel file, including formulas, into a dictionary of DataFrames.

    Args:
        file_path (str): Path to the Excel file.

    Returns:
        dict: A dictionary where keys are sheet names and values are DataFrames containing data and formulas.
    """
    try:
        # Load workbook without evaluating formulas
        wb = openpyxl.load_workbook(file_path, data_only=False)  # data_only=False to capture formulas
        sheet_names = wb.sheetnames
        all_dfs = {}

        # Iterate through each sheet
        for sheet_name in sheet_names:
            ws = wb[sheet_name]

            # Dynamically calculate the used range
            range_string = ws.calculate_dimension()
            min_col, min_row, max_col, max_row = range_boundaries(range_string)

            data_rows = []
            header_row = None

            # Iterate over the entire range to capture values and formulas
            for row_idx, row in enumerate(ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col), start=1):
                row_data = []
                print([str(cell.value) for cell in row]) 
                print([cell.data_type for cell in row])  # Check cell types
                for cell in row:
                    print(f"{cell.col_idx}-{cell.internal_value}")
                    # Check if cell contains a formula
                    if cell.data_type == 'f':  # Formula cell
                        # row_data.append(f"={cell.value}")
                        row_data.append(str(cell.internal_value))
                    else:
                        if row_idx == 1 and cell.value == None:
                            cell.value = ""
                        row_data.append(cell.value)

                # Store header row separately
                if row_idx == 1:
                    header_row = row_data
                else:
                    data_rows.append(row_data)

            # Create DataFrame if valid data is present
            if header_row and data_rows:
                df = pd.DataFrame(data_rows, columns=header_row)
                print(df)
                all_dfs[sheet_name] = df
            else:
                all_dfs[sheet_name] = pd.DataFrame()  # Empty DataFrame for empty sheets

        return all_dfs

    except FileNotFoundError:
        print(f"Error: File not found at {file_path}")
        return None
    except Exception as e:
        print(f"An error occurred: {e}")
        return None
    
def dataframe_to_excel_with_dataframe_formulas(df, output_file, sheet_name="Sheet1", formula_column="formula"):
    """
    Writes a pandas DataFrame to Excel, handling formulas stored within the DataFrame, and keeping the headers.

    Args:
        df (pd.DataFrame): The DataFrame to write.
        output_file (str): The path to the output Excel file.
        sheet_name (str): The name of the sheet.
        formula_column (str): The name of the column containing formulas (or None if no such column exists).
    """
    try:
        # Create a copy of the dataframe so as not to modify the original.
        df_copy = df.copy()

        # Extract formulas and remove formula column from the copy.
        formula_cells = {}
        if formula_column in df_copy.columns:
            for index, row in df_copy.iterrows():
                formula = row[formula_column]
                if pd.notna(formula) and isinstance(formula, str) and formula.startswith("="):
                    col_letter = openpyxl.utils.get_column_letter(df_copy.columns.get_loc(formula_column))
                    cell_address = f"{openpyxl.utils.get_column_letter(df_copy.columns.get_loc(formula_column)+1)}{index + 2}"
                    formula_cells[cell_address] = formula
            df_copy = df_copy.drop(formula_column, axis=1)

        # Write the DataFrame to Excel (without formula column), keeping headers
        with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
            df_copy.to_excel(writer, sheet_name=sheet_name, index=False) #Now write the headers.

        # Load the workbook and add formulas
        workbook = openpyxl.load_workbook(output_file)
        sheet = workbook[sheet_name]
        for cell_address, formula in formula_cells.items():
            sheet[cell_address] = formula

        # Save the workbook
        workbook.save(output_file)

        print(f"DataFrame written to '{output_file}' with DataFrame formulas and headers.")

    except Exception as e:
        print(f"An error occurred: {e}")

# Example usage:
file_path = r"C:\Users\gaurav.j.choubey\Desktop\project\gtic-7-2025\excelsheet_poc\excel_genai_poc\examplemacro1.xlsm"  # Replace with your file path

result_dfs = read_excel_with_formulas_all_sheets(file_path)
# result_dfs = pd.read_excel(file_path,sheet_name=None)

print(result_dfs)

# for sheet_name, df in result_dfs.items():
dataframe_to_excel_with_dataframe_formulas(df =result_dfs['Data'], output_file="./tmp.xlsx",sheet_name='Data')

# print(f"type{type(result_dfs)}")

# result_dfs.to_excel()



# if result_dfs:
    # for sheet_name, df in result_dfs.items():
    #     print(f"Sheet: {sheet_name}")
    #     if not df.empty:
    #         print(df)
    #     else:
    #         print("Sheet is empty or contains no valid data.")
    #     print("-" * 40)
